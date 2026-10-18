#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把 tools/_out 的 JSON 生成离线 Excel 数据表（研究分析用）。
用法：先 node tools/export-json.js，再 python tools/build-xlsx.py
"""
import json
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'tools', '_out')
DEST = os.path.join(ROOT, '鲛人研究数据库_条目表_2026-10.xlsx')

with open(os.path.join(OUT, 'entries.json'), encoding='utf-8') as f:
    entries = json.load(f)
with open(os.path.join(OUT, 'bibliography.json'), encoding='utf-8') as f:
    bib = json.load(f)
with open(os.path.join(OUT, 'schema.json'), encoding='utf-8') as f:
    schema = json.load(f)

TRAITS = schema['traits']
TRAIT_BY_ID = {t['id']: t for t in TRAITS}
KIND = schema['kindById']
CULTURES = [c['id'] for c in schema['cultures']]
CONFIDENCES = [c['id'] for c in schema['confidences']]

INK = '16221F'
TEAL = '1D5C54'
TEAL_SOFT = 'D7E2DA'
PAPER = 'EEF0E9'
SEAL = 'A63A2B'
LINE = 'C7D0C5'

HEADER_FONT = Font(name='等线', bold=True, color='FFFFFF', size=10)
HEADER_FILL = PatternFill('solid', fgColor=TEAL)
TITLE_FONT = Font(name='等线', bold=True, size=14, color=INK)
BODY_FONT = Font(name='等线', size=10, color=INK)
SMALL_FONT = Font(name='等线', size=9, color='66746E')
THIN = Side(style='thin', color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def trait_names(ids):
    return '、'.join(TRAIT_BY_ID[i]['label'] for i in ids if i in TRAIT_BY_ID)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def write_table(ws, headers, rows, widths, start_row=1):
    for i, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=i, value=h)
    style_header(ws, start_row, len(headers))
    for r, row in enumerate(rows, start_row + 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = 'A%d:%s%d' % (start_row, get_column_letter(len(headers)), start_row + len(rows))


wb = Workbook()

# ---------------------------------------------------------------- 条目总表
ws = wb.active
ws.title = '条目总表'
headers = ['编号', '类型', '子类', '条目标题', '作品/出处', '作者/创作者', '时代', '年代（年）',
           '年代显示', '谱系', '地区', '媒介', '主名', '异名', '原文词', '特征', '叙事功能',
           '标签', '原文引文', '引文出处', '今译/译名', '条目释读', '研究意义', '置信度',
           '出处列表', '关联条目']
rows = []
for e in entries:
    term = e.get('term') or {}
    sources = '\n'.join('%s（%s）' % (s.get('label', ''), s.get('detail', '')) for s in (e.get('sources') or []))
    rows.append([
        e.get('id', ''), KIND.get(e.get('kind'), e.get('kind', '')), e.get('subtype', ''),
        e.get('title', ''), e.get('work', ''), e.get('creator', ''), e.get('eraLabel', ''),
        e.get('year'), e.get('dateDisplay', ''), e.get('culture', ''), e.get('region', ''),
        e.get('medium', '') or '', term.get('primary', ''), '、'.join(term.get('variants') or []),
        term.get('original', ''), trait_names(e.get('traits') or []), trait_names(e.get('narratives') or []),
        '、'.join(e.get('tags') or []), e.get('excerpt', ''), e.get('excerptSource', ''),
        e.get('translation', ''), e.get('summary', ''), e.get('significance', ''), e.get('confidence', ''),
        sources, '、'.join(e.get('related') or []),
    ])
write_table(ws, headers, rows,
            [24, 14, 10, 30, 26, 18, 13, 9, 17, 8, 10, 8, 12, 16, 16, 28, 24, 14, 50, 26, 26, 55, 45, 9, 45, 26])
ws.row_dimensions[1].height = 26

# ---------------------------------------------------------------- 特征编码矩阵
ws2 = wb.create_sheet('特征编码')
headers2 = ['编号', '条目标题'] + [('%s·%s' % (t['group'], t['label'])) for t in TRAITS]
rows2 = []
for e in entries:
    ids = set((e.get('traits') or []) + (e.get('narratives') or []))
    rows2.append([e.get('id', ''), e.get('title', '')] + [1 if t['id'] in ids else 0 for t in TRAITS])
write_table(ws2, headers2, rows2, [24, 30] + [11] * len(TRAITS))
for c in range(3, len(headers2) + 1):
    for r in range(2, len(rows2) + 2):
        ws2.cell(row=r, column=c).alignment = CENTER
        ws2.cell(row=r, column=c).font = BODY_FONT
ws2.row_dimensions[1].height = 52

# ---------------------------------------------------------------- 名称与译名
ws3 = wb.create_sheet('名称与译名')
rows3 = []
for e in entries:
    term = e.get('term') or {}
    rows3.append([e.get('id', ''), e.get('title', ''), term.get('primary', ''),
                  '、'.join(term.get('variants') or []), term.get('original', ''),
                  term.get('note', ''), e.get('dateDisplay', ''), e.get('culture', '')])
write_table(ws3, ['编号', '条目标题', '主名', '异名', '原文词', '名称说明', '年代', '谱系'],
            rows3, [24, 32, 14, 26, 30, 55, 18, 9])

# ---------------------------------------------------------------- 参考书目
ws4 = wb.create_sheet('参考书目')
rows4 = [[b.get('category', ''), b.get('author', ''), b.get('title', ''), b.get('year', ''),
          b.get('publisher', ''), b.get('note', ''), b.get('url', '')] for b in bib]
write_table(ws4, ['分类', '作者/编者', '书名', '年份', '出版社', '说明', '链接'],
            rows4, [20, 26, 36, 8, 24, 60, 40])

# ---------------------------------------------------------------- 数据说明
ws5 = wb.create_sheet('数据说明')
ws5.column_dimensions['A'].width = 26
ws5.column_dimensions['B'].width = 92
years = [e.get('year') for e in entries if isinstance(e.get('year'), (int, float))]
year_range = '至'.join(str(y) for y in [min(years), max(years)]) + ' 年'
lines = [
    ('项目全称', '从鲛人到美人鱼——关于神话形象的跨文化重构的可视化研究'),
    ('数据版本', 'v' + str(schema.get('version', '1.0')) + '（2026 年 10 月）'),
    ('导出日期', date.today().isoformat()),
    ('条目总数', len(entries)),
    ('类型分布', ' / '.join('%s %d' % (KIND.get(k, k), sum(1 for e in entries if e.get('kind') == k))
                          for k in ['ancient', 'translation', 'modern'])),
    ('谱系分布', ' / '.join('%s %d' % (c, sum(1 for e in entries if e.get('culture') == c)) for c in CULTURES)),
    ('置信度', ' / '.join('%s %d' % (c, sum(1 for e in entries if e.get('confidence') == c)) for c in CONFIDENCES)),
    ('年代范围', year_range),
    ('编码维度', '%d 项（形态 9 / 能力 17 / 叙事 13）' % len(TRAITS)),
    ('', ''),
    ('使用说明', '本表为研究辅助数据，不是定本。引用前请核对条目 sources 与 confidence；「待核」条目的线索与出处须回原书复核。'),
    ('字段对照', '「特征编码」表用 1/0 表示条目是否具备该特征；1 只表示已编码，不表示文本中必然存在该特征。'),
    ('版权说明', '当代作品仅作简短引用与信息性著录；古籍引文属研究性引用。'),
    ('网站', 'index.html（纯静态，可离线打开）；方法文档见 docs/ 目录。'),
]
ws5['A1'] = '鲛人研究数据库 · 数据说明'
ws5['A1'].font = TITLE_FONT
ws5.merge_cells('A1:B1')
for i, (k, v) in enumerate(lines, 3):
    ws5.cell(row=i, column=1, value=k).font = Font(name='等线', bold=True, size=10, color=TEAL)
    cell = ws5.cell(row=i, column=2, value=v)
    cell.font = BODY_FONT
    cell.alignment = WRAP

# ---------------------------------------------------------------- 样式收尾
for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows():
        for cell in row:
            if cell.row == 1 and sheet.title != '数据说明':
                continue
            if cell.value is not None and cell.font is None:
                cell.font = BODY_FONT

wb.save(DEST)
print('已生成：' + DEST)
print('sheets:', ', '.join(ws.title for ws in wb.worksheets))
